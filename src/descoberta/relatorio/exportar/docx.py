"""Word (DOCX) com python-docx: não precisa do Word instalado."""

import io
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..documento import (
    Assinatura,
    Aviso,
    Capa,
    Lista,
    Paragrafo,
    QuebraDePagina,
    Tabela,
    Titulo,
)
from .estilo import COR_AVISO, COR_CABECALHO, COR_DESTAQUE, CORES_NIVEL

if TYPE_CHECKING:
    from ..emissao import RelatorioEmitido

LARGURA_UTIL_CM = 17.0  # A4 (21 cm) menos as margens de 2 cm


def _sombrear(elemento, cor: str) -> None:
    """Cor de fundo de uma célula ou parágrafo (o python-docx não tem isso pronto)."""
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor.lstrip("#"))
    propriedades = (
        elemento.get_or_add_tcPr()
        if hasattr(elemento, "get_or_add_tcPr")
        else (elemento.get_or_add_pPr())
    )
    propriedades.append(sombra)


def _campo(paragrafo, codigo: str) -> None:
    """Campo do Word (PAGE, NUMPAGES): o Word calcula o número ao abrir o arquivo."""
    run = paragrafo.add_run()
    for tipo, texto in (("begin", None), (None, codigo), ("end", None)):
        if tipo:
            marca = OxmlElement("w:fldChar")
            marca.set(qn("w:fldCharType"), tipo)
            run._r.append(marca)
        else:
            instrucao = OxmlElement("w:instrText")
            instrucao.set(qn("xml:space"), "preserve")
            instrucao.text = texto
            run._r.append(instrucao)


def _repetir_cabecalho(linha) -> None:
    propriedades = linha._tr.get_or_add_trPr()
    cabecalho = OxmlElement("w:tblHeader")
    cabecalho.set(qn("w:val"), "true")
    propriedades.append(cabecalho)


def _tabela(doc, tabela: Tabela) -> None:
    larguras = tabela.larguras or [1] * len(tabela.colunas)
    total = sum(larguras)
    grade = doc.add_table(rows=1 + len(tabela.linhas), cols=len(tabela.colunas))
    grade.style = "Table Grid"
    for indice, linha in enumerate([tabela.colunas, *tabela.linhas]):
        for coluna, texto in enumerate(linha):
            celula = grade.rows[indice].cells[coluna]
            celula.width = Cm(LARGURA_UTIL_CM * larguras[coluna] / total)
            celula.text = ""
            partes = texto.split("\n")
            run = celula.paragraphs[0].add_run(partes[0])
            for parte in partes[1:]:
                run.add_break()
                run.add_text(parte)
            run.font.size = Pt(9)
            if indice == 0:
                run.bold = True
                _sombrear(celula._tc, COR_CABECALHO)
            elif tabela.niveis and (nivel := tabela.niveis[indice - 1][coluna]):
                _sombrear(celula._tc, CORES_NIVEL[nivel])
    _repetir_cabecalho(grade.rows[0])
    doc.add_paragraph()


def _bloco(doc, bloco) -> None:
    match bloco:
        case Capa():
            if bloco.logotipo:
                doc.add_picture(io.BytesIO(bloco.logotipo), height=Cm(2.5))
            if bloco.faixa:
                faixa = doc.add_paragraph()
                run = faixa.add_run(bloco.faixa)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(COR_DESTAQUE.lstrip("#"))
            doc.add_paragraph(bloco.titulo, style="Title")
            subtitulo = doc.add_paragraph()
            subtitulo.add_run(bloco.subtitulo).font.size = Pt(18)
            for rotulo, valor in bloco.linhas:
                paragrafo = doc.add_paragraph()
                paragrafo.add_run(f"{rotulo}: ").bold = True
                paragrafo.add_run(valor)
            if bloco.nota:
                doc.add_paragraph().add_run(bloco.nota).italic = True
        case Titulo():
            titulo = doc.add_heading(bloco.texto, level=bloco.nivel)
            titulo.paragraph_format.keep_with_next = True  # nunca sozinho no pé da página
        case Paragrafo():
            paragrafo = doc.add_paragraph()
            if bloco.rotulo:
                paragrafo.add_run(f"{bloco.rotulo} ").bold = True
            paragrafo.add_run(bloco.texto)
        case Aviso():
            paragrafo = doc.add_paragraph()
            paragrafo.add_run(bloco.texto).bold = True
            _sombrear(paragrafo._p, COR_AVISO)
        case Lista():
            for item in bloco.itens:
                doc.add_paragraph(item, style="List Bullet")
        case Tabela():
            _tabela(doc, bloco)
        case Assinatura():
            titulo = doc.add_paragraph()
            titulo.paragraph_format.space_before = Pt(18)
            titulo.add_run(bloco.titulo).bold = True
            for campo in bloco.campos:
                linha = doc.add_paragraph()
                linha.paragraph_format.space_before = Pt(14 if campo == "Assinatura" else 8)
                linha.add_run(f"{campo}: ").bold = True
                linha.add_run(bloco.valor(campo) or "_" * 50)
        case QuebraDePagina():
            doc.add_page_break()
        case _:
            raise TypeError(f"bloco desconhecido: {bloco!r}")


def exportar_docx(emitido: "RelatorioEmitido") -> bytes:
    doc = Document()
    doc.core_properties.title = emitido.documento.titulo
    doc.styles["Normal"].font.size = Pt(10)
    secao = doc.sections[0]
    secao.page_height, secao.page_width = Cm(29.7), Cm(21)
    for margem in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(secao, margem, Cm(2))
    rodape = secao.footer.paragraphs[0]
    rodape.text = f"{emitido.rodape} · Página "
    _campo(rodape, "PAGE")
    rodape.add_run(" de ")
    _campo(rodape, "NUMPAGES")
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in rodape.runs:
        run.font.size = Pt(7)
    for bloco in emitido.documento.blocos:
        _bloco(doc, bloco)
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
