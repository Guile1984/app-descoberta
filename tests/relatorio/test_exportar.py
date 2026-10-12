"""Exportação para Markdown, DOCX e PDF; hash no rodapé; geração sem internet."""

import hashlib
import io
import socket
from datetime import UTC, datetime

import pytest
from docx import Document as AbrirDocx
from pypdf import PdfReader

from descoberta.relatorio import (
    Emissao,
    emitir,
    exportar_docx,
    exportar_markdown,
    exportar_pdf,
    montar_relatorio,
)
from descoberta.relatorio.documento import Documento, Paragrafo, Tabela
from descoberta.relatorio.exportar.markdown import markdown_do_documento
from descoberta.repositorio import TipoVersao

EMITIDO_EM = datetime(2026, 10, 11, 17, 5, tzinfo=UTC)


def _texto_do_pdf(dados: bytes) -> str:
    leitor = PdfReader(io.BytesIO(dados))
    return "\n".join(pagina.extract_text() for pagina in leitor.pages)


def _texto_do_docx(dados: bytes) -> str:
    documento = AbrirDocx(io.BytesIO(dados))
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        partes += [celula.text for linha in tabela.rows for celula in linha.cells]
    return "\n".join(partes)


# --- Emissão e hash ---


def test_hash_e_o_sha256_do_markdown_canonico(dados_roberto, emissao_cliente):
    emitido = emitir(dados_roberto, emissao_cliente)
    esperado = hashlib.sha256(emitido.markdown.encode("utf-8")).hexdigest()
    assert emitido.hash_conteudo == esperado
    assert emitido.markdown == markdown_do_documento(emitido.documento)


def test_mesmos_dados_mesmo_hash(dados_roberto, emissao_cliente):
    assert emitir(dados_roberto, emissao_cliente).hash_conteudo == (
        emitir(dados_roberto, emissao_cliente).hash_conteudo
    )


def test_numero_e_tipo_mudam_o_hash(dados_roberto, emissao_cliente, emissao_interna):
    hashes = {
        emitir(dados_roberto, emissao_cliente).hash_conteudo,
        emitir(dados_roberto, emissao_interna).hash_conteudo,
        emitir(dados_roberto, Emissao(TipoVersao.cliente, 1, EMITIDO_EM)).hash_conteudo,
    }
    assert len(hashes) == 3


def test_rodape_de_versao_registrada(dados_roberto):
    emitido = emitir(dados_roberto, Emissao(TipoVersao.cliente, 3, EMITIDO_EM))
    assert "Versão 3 (cliente)" in emitido.rodape
    assert "11/10/2026" in emitido.rodape
    assert f"hash {emitido.hash_conteudo[:12]}" in emitido.rodape


def test_rodape_de_rascunho(dados_roberto, emissao_interna):
    rodape = emitir(dados_roberto, emissao_interna).rodape
    assert "Rascunho não registrado (interna)" in rodape


# --- Markdown ---


def test_markdown_do_roberto(dados_roberto, emissao_cliente):
    emitido = emitir(dados_roberto, emissao_cliente)
    texto = exportar_markdown(emitido)
    assert texto.startswith("# Relatório de Descoberta")
    assert "## 1. Resumo executivo" in texto
    assert "| Indicador | Camada | Valor atual | Meta | Prazo | Como medir |" in texto
    assert "**Nome:** ____" in texto
    assert texto.rstrip().endswith(emitido.rodape)  # rodapé no fim, fora do hash


def test_markdown_escapa_tabelas_e_marcacao():
    documento = Documento(
        "Teste",
        [
            Paragrafo("Use *sempre* o modelo #1 <b>"),
            Tabela(["A", "B"], [["x | y", "linha 1\nlinha 2"]]),
        ],
    )
    texto = markdown_do_documento(documento)
    assert r"Use \*sempre\* o modelo \#1 \<b\>" in texto
    assert r"| x \| y | linha 1<br>linha 2 |" in texto


# --- DOCX ---


def test_docx_do_roberto(dados_roberto, emissao_cliente):
    emitido = emitir(dados_roberto, Emissao(TipoVersao.cliente, 2, EMITIDO_EM))
    dados = exportar_docx(emitido)
    texto = _texto_do_docx(dados)
    assert "Relatório de Descoberta" in texto
    assert "12. Recomendação e próximos passos" in texto
    assert "Termo de ciência e validação" in texto
    assert "Recall de fissuras graves" in texto
    documento = AbrirDocx(io.BytesIO(dados))
    rodape = documento.sections[0].footer.paragraphs[0].text
    assert emitido.hash_conteudo[:12] in rodape
    assert "Versão 2 (cliente)" in rodape
    titulos = [p.text for p in documento.paragraphs if p.style.name == "Heading 1"]
    assert titulos[0] == "1. Resumo executivo"


# --- PDF ---


def test_pdf_do_roberto(dados_roberto):
    emitido = emitir(dados_roberto, Emissao(TipoVersao.cliente, 1, EMITIDO_EM))
    dados = exportar_pdf(emitido)
    assert dados.startswith(b"%PDF")
    leitor = PdfReader(io.BytesIO(dados))
    assert len(leitor.pages) >= 5
    texto = _texto_do_pdf(dados)
    assert "Construtora Alvorada" in texto
    assert "Termo de ciência e validação" in texto
    assert "≥ 95%" in texto  # fonte DejaVu: caracteres fora do Latin-1
    assert "não é proposta comercial nem contrato" in texto
    # Rodapé em todas as páginas, com o hash e o número da página.
    for numero, pagina in enumerate(leitor.pages, start=1):
        texto_pagina = pagina.extract_text()
        assert emitido.hash_conteudo[:12] in texto_pagina
        assert f"Página {numero} de {len(leitor.pages)}" in texto_pagina


def test_pdf_da_versao_interna_avisa_em_todas_as_paginas(dados_roberto, emissao_interna):
    dados = exportar_pdf(emitir(dados_roberto, emissao_interna))
    for pagina in PdfReader(io.BytesIO(dados)).pages:
        assert "USO INTERNO" in pagina.extract_text()


def test_pdf_de_sessao_vazia(dados_vazios, emissao_cliente):
    assert exportar_pdf(emitir(dados_vazios, emissao_cliente)).startswith(b"%PDF")


# --- Critério de aceite: relatório do Roberto gerado sem internet ---


@pytest.fixture
def sem_rede(monkeypatch):
    def recusar(*args, **kwargs):
        raise OSError("rede bloqueada no teste")

    monkeypatch.setattr(socket, "socket", recusar)
    monkeypatch.setattr(socket, "create_connection", recusar)
    monkeypatch.setattr(socket, "getaddrinfo", recusar)


@pytest.mark.parametrize("versao", [TipoVersao.cliente, TipoVersao.interna])
def test_relatorio_do_roberto_offline_nos_tres_formatos(dados_roberto, sem_rede, versao):
    emitido = emitir(dados_roberto, Emissao(versao, 1, EMITIDO_EM))
    assert "12 de 18" in exportar_markdown(emitido)
    assert "12 de 18" in _texto_do_docx(exportar_docx(emitido))
    assert "12 de 18" in _texto_do_pdf(exportar_pdf(emitido))


def test_montagem_nao_muda_os_dados(dados_roberto, emissao_cliente):
    antes = dict(dados_roberto.respostas)
    montar_relatorio(dados_roberto, emissao_cliente)
    assert dados_roberto.respostas == antes


@pytest.mark.parametrize("tipo", list(TipoVersao))
def test_nenhuma_palavra_quebra_nas_tabelas_do_pdf(dados_roberto, tipo):
    """Cada palavra das tabelas cabe na coluna ("Probabilidad e" não pode acontecer)."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from descoberta.relatorio.exportar import pdf

    pdf._registrar_fontes()
    documento = emitir(dados_roberto, Emissao(tipo, None, EMITIDO_EM)).documento
    quebradas = []
    for tabela in (b for b in documento.blocos if isinstance(b, Tabela)):
        larguras = tabela.larguras or [1] * len(tabela.colunas)
        for linha, fonte in [(tabela.colunas, pdf.FONTE_NEGRITO)] + [
            (linha, pdf.FONTE) for linha in tabela.linhas
        ]:
            for celula, parte in zip(linha, larguras, strict=True):
                util = pdf.LARGURA_UTIL * parte / sum(larguras) - 12  # 6 pt de cada lado
                quebradas += [
                    (tabela.colunas[0], palavra)
                    for palavra in celula.split()
                    if stringWidth(palavra, fonte, 8.5) > util
                ]
    assert quebradas == []


# --- Logotipo e assinatura do consultor ---


def test_logotipo_na_capa_do_pdf_e_do_word(dados_roberto, emissao_cliente):
    emitido = emitir(dados_roberto, emissao_cliente)
    primeira = PdfReader(io.BytesIO(exportar_pdf(emitido))).pages[0]
    assert len(primeira.images) == 1
    assert len(AbrirDocx(io.BytesIO(exportar_docx(emitido))).inline_shapes) == 1


def test_sem_logotipo_nao_ha_imagem(dados_vazios, emissao_cliente):
    emitido = emitir(dados_vazios, emissao_cliente)
    assert not PdfReader(io.BytesIO(exportar_pdf(emitido))).pages[0].images
    assert len(AbrirDocx(io.BytesIO(exportar_docx(emitido))).inline_shapes) == 0


def test_assinatura_do_consultor_preenchida_nos_tres_formatos(dados_roberto, emissao_cliente):
    emitido = emitir(dados_roberto, emissao_cliente)
    for texto in (
        exportar_markdown(emitido),
        _texto_do_pdf(exportar_pdf(emitido)),
        _texto_do_docx(exportar_docx(emitido)),
    ):
        assert "Ana Exemplo" in texto
        assert "11.222.333/0001-81" in texto
